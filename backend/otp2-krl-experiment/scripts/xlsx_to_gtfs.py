#!/usr/bin/env python3
"""Task 1: arah sebaliknya dari gtfs_to_xlsx.py — xlsx (hasil suntingan Bayu di
Google Sheets) kembali menjadi file GTFS, supaya bisa divalidasi ulang.

    xlsx_to_gtfs.py <input.xlsx> <folder_output_atau_zip>

Tab yang namanya bukan nama file GTFS (mis. "Ringkasan", "Baca Saya") dilewati
dengan pesan, bukan diam-diam dibuang — sheet Bayu memang memuat tab catatan
seperti itu.
"""
import csv
import sys
import zipfile
from pathlib import Path

import openpyxl

# Tab di luar daftar ini dianggap catatan manusia, bukan tabel GTFS.
GTFS_FILES = {
    "agency", "stops", "routes", "trips", "stop_times", "calendar",
    "calendar_dates", "fare_attributes", "fare_rules", "transfers", "shapes",
    "frequencies", "feed_info", "ticketing_identifiers", "ticketing_deep_links",
}


def cell(value) -> str:
    """Sel kosong -> string kosong, bukan "None"."""
    return "" if value is None else str(value)


def main() -> None:
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    src, out = Path(sys.argv[1]), Path(sys.argv[2])
    if not src.exists():
        sys.exit(f"ERROR: {src} tidak ada")

    wb = openpyxl.load_workbook(src, read_only=True)
    tables: dict[str, tuple[list[str], list[list[str]]]] = {}

    for name in wb.sheetnames:
        if name not in GTFS_FILES:
            print(f"  dilewati (bukan tabel GTFS): {name}")
            continue
        rows = wb[name].iter_rows(values_only=True)
        header = next(rows, None)
        if header is None:
            print(f"  dilewati (kosong): {name}")
            continue
        header = [cell(h) for h in header if h is not None]
        body = [
            [cell(v) for v in row[: len(header)]]
            for row in rows
            if any(v is not None and str(v) != "" for v in row)
        ]
        tables[name] = (header, body)
        print(f"  {name:<24}{len(body):>8} baris x {len(header):>2} kolom")
    wb.close()

    if not tables:
        sys.exit("ERROR: tidak ada tab GTFS yang bisa dibaca")

    def write_csv(fh, header, body) -> None:
        # newline="" + lineterminator eksplisit: tanpa ini Windows menulis CRLF
        # ganda, dan file hasilnya tidak lagi identik dengan feed aslinya.
        writer = csv.writer(fh, lineterminator="\n")
        writer.writerow(header)
        writer.writerows(body)

    if out.suffix == ".zip":
        out.parent.mkdir(parents=True, exist_ok=True)
        with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zf:
            for name, (header, body) in tables.items():
                buf = []

                class _Sink:
                    def write(self, s):
                        buf.append(s)

                write_csv(_Sink(), header, body)
                zf.writestr(f"{name}.txt", "".join(buf))
    else:
        out.mkdir(parents=True, exist_ok=True)
        for name, (header, body) in tables.items():
            with (out / f"{name}.txt").open("w", encoding="utf-8", newline="") as fh:
                write_csv(fh, header, body)

    print(f"\n  {len(tables)} file GTFS -> {out}")


if __name__ == "__main__":
    main()
